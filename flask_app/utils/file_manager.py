import os
import pandas as pd
from datetime import datetime
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo
from flask import send_file # Para la función de descarga
from pathlib import Path
import base64
import uuid

import cv2
import numpy as np

# Definimos las rutas de los archivos de texto
BASE_DIR = os.path.dirname(os.path.abspath(__file__))
DATA_DIR = os.path.abspath(os.path.join(BASE_DIR, '..', 'data'))

# Asegúrate de que la carpeta 'data' exista
if not os.path.exists(DATA_DIR):
    os.makedirs(DATA_DIR)

ARCHIVO_PERSONAL_SONDA = os.path.join(DATA_DIR, "PERSONAL_SONDA.txt")
ARCHIVO_PERSONAL_EXTERNO = os.path.join(DATA_DIR, "PERSONAL_EXTERNO.txt")
ARCHIVO_MODELOS_RADIO = os.path.join(DATA_DIR, "MODELOS_RADIO.txt")
ARCHIVO_EMPRESAS_EXTERNAS = os.path.join(DATA_DIR, "EMPRESAS_EXTERNAS.txt")
ARCHIVO_AUTORIZADO_POR = os.path.join(DATA_DIR, "AUTORIZADO_POR.txt")

# Ruta del archivo Excel para el libro de radio
ARCHIVO_LIBRO_RADIO_EXCEL = os.path.join(DATA_DIR, "libro_radio.xlsx")

EXCEL_TABLE_NAME = "Tabla1" 

FACE_DATA_DIR = os.path.join(DATA_DIR, "faces")
FACE_EXCEL_FILE = os.path.join(DATA_DIR, "registro_accesos.xlsx")
FACE_TXT_FILE = os.path.join(DATA_DIR, "registro_accesos.txt")
FACE_REGISTRY_SHEET = "Rostros"
FACE_REGISTRY_TABLE = "TablaRostros"

FACE_LOG_SHEET = "Accesos"
FACE_LOG_TABLE = "TablaAccesos"

FACE_HEADERS = ["ID", "Nombre", "Archivo", "Fecha Registro"]
FACE_LOG_HEADERS = ["ID", "Nombre", "Tipo", "Fecha", "Hora", "Timestamp"]

CASCADE_PATH = cv2.data.haarcascades + 'haarcascade_frontalface_default.xml'

ARCHIVOS_DATA = {
    "personal_sonda": ARCHIVO_PERSONAL_SONDA,
    "personal_externo": ARCHIVO_PERSONAL_EXTERNO,
    "modelos_radio": ARCHIVO_MODELOS_RADIO,
    "empresas_externas": ARCHIVO_EMPRESAS_EXTERNAS,
    "autorizado_por": ARCHIVO_AUTORIZADO_POR,
}

def leer_datos(nombre_archivo):
    """Lee líneas de un archivo de texto y las devuelve como una lista."""
    ruta_archivo = ARCHIVOS_DATA.get(nombre_archivo)
    if not ruta_archivo or not os.path.exists(ruta_archivo):
        return []
    with open(ruta_archivo, 'r', encoding='utf-8') as f:
        datos = [line.strip() for line in f if line.strip()]
    return datos

def escribir_datos(nombre_archivo, datos_lista):
    """Escribe una lista de cadenas en un archivo de texto, una por línea."""
    ruta_archivo = ARCHIVOS_DATA.get(nombre_archivo)
    if not ruta_archivo:
        print(f"Error: Archivo '{nombre_archivo}' no reconocido.")
        return False
    try:
        with open(ruta_archivo, 'w', encoding='utf-8') as f:
            for item in datos_lista:
                f.write(f"{item}\n")
        return True
    except Exception as e:
        print(f"Error al escribir en {ruta_archivo}: {e}")
        return False

def agregar_dato(nombre_archivo, nuevo_dato):
    """Agrega un nuevo dato al archivo si no existe ya."""
    datos = leer_datos(nombre_archivo)
    if nuevo_dato not in datos:
        datos.append(nuevo_dato)
        return escribir_datos(nombre_archivo, datos)
    return False

def eliminar_dato(nombre_archivo, dato_a_eliminar):
    """Elimina un dato de un archivo."""
    datos = leer_datos(nombre_archivo)
    if dato_a_eliminar in datos:
        datos.remove(dato_a_eliminar)
        return escribir_datos(nombre_archivo, datos)
    return False

# --- MAPEO DE COLUMNAS PARA EXCEL (para openpyxl) ---
# Revertido a la estructura original sin las nuevas columnas
EXCEL_COLUMN_MAP = {
    'ID': 1, # Columna A
    'OT': 2, # Columna B
    'Fecha Recepción': 3, # Columna C
    'Personal Sonda Receptor': 4, # Columna D
    'Personal Externo Entrega': 5, # Columna E
    'Tipo de Trabajo': 6, # Columna F
    'Empresa': 7, # Columna G
    'Nombre Autorizado': 8, # Columna H
    'Área': 9, # Columna I
    'Visto Bueno': 10, # Columna J
    'Dato Empresa': 11, # Columna K
    'Factura': 12, # Columna L
    'ID Empresa': 13, # Columna M
    'Solicitud O/D': 14, # Columna N
    'Solicitud Anexo': 15, # Columna O
    'Modelo de Radio': 16, # Columna P
    'Cantidad de Radios': 17, # Columna Q
    'Programador Sonda': 18, # Columna R
    'Fecha de Entrega Sonda': 19, # Columna S
    'Personal Sonda (Entrega)': 20, # Columna T
    'Personal que Recibe': 21, # Columna U
    'Estado': 22 # Columna V
}

# Fila donde comienzan los encabezados en tu Excel (según la imagen)
HEADER_ROW_EXCEL = 3
# Fila donde comienzan los datos en tu Excel (justo después de los encabezados)
DATA_START_ROW_EXCEL = HEADER_ROW_EXCEL + 1

# --- FUNCIONES PARA MANEJAR EL EXCEL CON OpenPyXL (CON RECREACIÓN DE TABLA) ---

def _recreate_excel_table(ws):
    """
    Elimina la tabla existente y la recrea con el rango actualizado.
    """
    if EXCEL_TABLE_NAME in ws.tables:
        del ws.tables[EXCEL_TABLE_NAME]
        print(f"DEBUG: Tabla '{EXCEL_TABLE_NAME}' eliminada para recreación.")

    # Asegurarse de que haya al menos una fila de encabezado para que la tabla sea válida
    if ws.max_row < HEADER_ROW_EXCEL:
        print(f"Advertencia: ws.max_row ({ws.max_row}) es menor que HEADER_ROW_EXCEL ({HEADER_ROW_EXCEL}). No se puede recrear la tabla.")
        return # No se puede recrear la tabla si no hay fila de encabezado

    last_col_letter = get_column_letter(max(EXCEL_COLUMN_MAP.values()))
    
    # El rango de la tabla va desde la primera columna (A) de la fila de encabezados (HEADER_ROW_EXCEL)
    # hasta la última columna de la última fila con datos (ws.max_row).
    full_table_range = f"{get_column_letter(1)}{HEADER_ROW_EXCEL}:{last_col_letter}{ws.max_row}"
    
    table_style_info = TableStyleInfo(name="TableStyleMedium9", showFirstColumn=False,
                                     showLastColumn=False, showRowStripes=True, showColumnStripes=False)
    new_table = Table(displayName=EXCEL_TABLE_NAME, ref=full_table_range)
    new_table.tableStyleInfo = table_style_info
    
    ws.add_table(new_table)
    print(f"DEBUG: Tabla '{EXCEL_TABLE_NAME}' recreada con rango: {full_table_range}")


def guardar_ot_en_excel(data):
    """
    Guarda una nueva OT en el archivo Excel 'libro_radio.xlsx'.
    Añade datos a partir de la fila 4 y recrea el objeto de tabla.
    """
    try:
        print(f"DEBUG: [guardar_ot_en_excel] Intentando guardar nueva OT: {data.get('OT', 'N/A')}")
        
        for key in ['Fecha Recepción', 'Fecha de Entrega Sonda']:
            if isinstance(data.get(key), str) and data.get(key):
                data[key] = datetime.strptime(data[key], '%d-%m-%Y').strftime('%d-%m-%Y')
            elif data.get(key) is None:
                data[key] = '' 

        if not os.path.exists(ARCHIVO_LIBRO_RADIO_EXCEL):
            print(f"DEBUG: [guardar_ot_en_excel] Creando nuevo archivo Excel: {ARCHIVO_LIBRO_RADIO_EXCEL}")
            wb = load_workbook() 
            ws = wb.active
            ws.title = "Libro de Radios"
            for col_name, col_idx in EXCEL_COLUMN_MAP.items():
                ws.cell(row=HEADER_ROW_EXCEL, column=col_idx, value=col_name) 
        else:
            print(f"DEBUG: [guardar_ot_en_excel] Cargando archivo Excel existente: {ARCHIVO_LIBRO_RADIO_EXCEL}")
            wb = load_workbook(ARCHIVO_LIBRO_RADIO_EXCEL)
            ws = wb.active 
            
        last_id = 0
        for row_idx in range(ws.max_row, DATA_START_ROW_EXCEL -1, -1): 
            cell_value = ws.cell(row=row_idx, column=EXCEL_COLUMN_MAP['ID']).value
            if cell_value is not None:
                try:
                    last_id = int(str(cell_value).strip())
                    break
                except (ValueError, TypeError):
                    pass 
        new_id = last_id + 1
        
        next_row = ws.max_row + 1 if ws.max_row >= HEADER_ROW_EXCEL else DATA_START_ROW_EXCEL
        
        print(f"DEBUG: [guardar_ot_en_excel] Próxima fila para escribir: {next_row}, Nuevo ID: {new_id}")

        data['ID'] = new_id
        
        max_col_index = max(EXCEL_COLUMN_MAP.values())
        row_values = [None] * max_col_index 
        
        for col_name, col_idx in EXCEL_COLUMN_MAP.items():
            value_to_write = data.get(col_name)
            if col_name in ['Visto Bueno', 'Dato Empresa', 'Factura', 'ID Empresa', 'Solicitud O/D', 'Solicitud Anexo']:
                value_to_write = "SI" if value_to_write == 'SI' or value_to_write == True else "NO"
            elif col_name == 'ID': # Asegurarse de que el ID se escriba como entero
                value_to_write = int(value_to_write) 
            elif value_to_write is None:
                value_to_write = '' 

            # Asegurarse de que el índice de la columna es válido
            if col_idx - 1 < len(row_values):
                row_values[col_idx - 1] = value_to_write
            else:
                print(f"Advertencia: Columna {col_name} (índice {col_idx}) está fuera del rango de columnas esperadas.")

        print(f"DEBUG: [guardar_ot_en_excel] Valores de fila a escribir: {row_values}")
        for col_idx, value in enumerate(row_values, start=1):
            ws.cell(row=next_row, column=col_idx, value=value)

        _recreate_excel_table(ws)

        wb.save(ARCHIVO_LIBRO_RADIO_EXCEL)
        print(f"DEBUG: [guardar_ot_en_excel] OT guardada exitosamente en: {ARCHIVO_LIBRO_RADIO_EXCEL}")
        return True
    except Exception as e:
        print(f"ERROR: [guardar_ot_en_excel] Fallo al guardar OT en Excel: {e}")
        return False


def leer_libro_radio_excel():
    """
    Lee todos los datos del archivo Excel 'libro_radio.xlsx' y los devuelve como una lista de diccionarios.
    Asume que los encabezados están en la fila 3 y los datos comienzan en la fila 4.
    """
    if not os.path.exists(ARCHIVO_LIBRO_RADIO_EXCEL):
        print(f"DEBUG: [leer_libro_radio_excel] Archivo Excel no encontrado: {ARCHIVO_LIBRO_RADIO_EXCEL}")
        return []
    
    try:
        print(f"DEBUG: [leer_libro_radio_excel] Leyendo archivo Excel: {ARCHIVO_LIBRO_RADIO_EXCEL}")
        wb = load_workbook(ARCHIVO_LIBRO_RADIO_EXCEL)
        ws = wb.active 

        data = []
        headers = [ws.cell(row=HEADER_ROW_EXCEL, column=col.column).value for col in ws[HEADER_ROW_EXCEL]]
        header_to_col_idx = {header: idx + 1 for idx, header in enumerate(headers) if header is not None}

        for expected_header in EXCEL_COLUMN_MAP.keys():
            if expected_header not in header_to_col_idx:
                print(f"Advertencia: Encabezado '{expected_header}' no encontrado en la fila {HEADER_ROW_EXCEL} del Excel.")
                header_to_col_idx[expected_header] = -1 


        for row_idx in range(DATA_START_ROW_EXCEL, ws.max_row + 1):
            row_dict = {}
            current_row_id = None 

            for col_name, col_map_idx in EXCEL_COLUMN_MAP.items():
                actual_col_idx = header_to_col_idx.get(col_name)
                
                if actual_col_idx is not None and actual_col_idx != -1: 
                    cell_value = ws.cell(row=row_idx, column=actual_col_idx).value
                    
                    if col_name == 'ID':
                        if cell_value is None:
                            current_row_id = None 
                        else:
                            try:
                                # Intenta convertir a string, luego limpia espacios, luego a entero
                                cleaned_value = str(cell_value).strip()
                                # Si el valor es una cadena vacía después de limpiar, trátalo como None
                                if cleaned_value: 
                                    current_row_id = int(cleaned_value)
                                else:
                                    current_row_id = None 
                            except (ValueError, TypeError):
                                print(f"Advertencia: [leer_libro_radio_excel] Fila {row_idx} tiene ID inválido '{cell_value}'. Estableciendo ID a None.")
                                current_row_id = None
                        row_dict['ID'] = current_row_id 
                    else:
                        row_dict[col_name] = cell_value if cell_value is not None else None
                else:
                    row_dict[col_name] = None 

            if row_dict.get('ID') is not None: 
                data.append(row_dict)
            else:
                print(f"Advertencia: Fila {row_idx} omitida porque su ID es nulo o inválido.")

        print(f"DEBUG: [leer_libro_radio_excel] Total de OTs leídas: {len(data)}")
        return data
    except Exception as e:
        print(f"ERROR: [leer_libro_radio_excel] Fallo al leer libro de radio desde Excel: {e}")
        return []


def actualizar_ot_en_excel(ot_id, updated_data):
    """
    Actualiza una OT existente en el archivo Excel 'libro_radio.xlsx'.
    Asume que los encabezados están en la fila 3 y los datos comienzan en la fila 4.
    """
    try:
        print(f"DEBUG: [actualizar_ot_en_excel] Intentando actualizar OT con ID: {ot_id}, Datos: {updated_data.get('OT', 'N/A')}")
        print(f"DEBUG: [actualizar_ot_en_excel] updated_data recibido: {updated_data}") 

        if not os.path.exists(ARCHIVO_LIBRO_RADIO_EXCEL):
            print(f"DEBUG: [actualizar_ot_en_excel] Archivo Excel no encontrado para actualización: {ARCHIVO_LIBRO_RADIO_EXCEL}")
            return False

        wb = load_workbook(ARCHIVO_LIBRO_RADIO_EXCEL)
        ws = wb.active 

        for key in ['Fecha Recepción', 'Fecha de Entrega Sonda']:
            if isinstance(updated_data.get(key), str) and updated_data.get(key):
                updated_data[key] = datetime.strptime(updated_data[key], '%d-%m-%Y').strftime('%d-%m-%Y')
            elif updated_data.get(key) is None:
                updated_data[key] = ''
        
        found_row = -1
        headers = [ws.cell(row=HEADER_ROW_EXCEL, column=col.column).value for col in ws[HEADER_ROW_EXCEL]]
        header_to_col_idx = {header: idx + 1 for idx, header in enumerate(headers) if header is not None}

        id_col_idx = header_to_col_idx.get('ID')
        if id_col_idx is None or id_col_idx == -1:
            print(f"ERROR: [actualizar_ot_en_excel] La columna 'ID' no se encontró en la fila de encabezados del Excel.")
            return False

        # Find the row to update
        for row_idx in range(DATA_START_ROW_EXCEL, ws.max_row + 1): 
            cell_value = ws.cell(row=row_idx, column=id_col_idx).value
            if cell_value is not None:
                try:
                    cleaned_value = str(cell_value).strip()
                    if cleaned_value and int(cleaned_value) == int(ot_id):
                        found_row = row_idx
                        break
                except (ValueError, TypeError):
                    pass 
        
        if found_row == -1:
            print(f"DEBUG: [actualizar_ot_en_excel] OT con ID {ot_id} no encontrada para actualización.")
            return False 

        print(f"DEBUG: [actualizar_ot_en_excel] OT con ID {ot_id} encontrada en fila {found_row}. Actualizando...")

        # Explicitly write the ID back to ensure it's not lost
        # Get the current ID value from the cell before overwriting other data
        current_id_in_excel_cell = ws.cell(row=found_row, column=id_col_idx).value
        print(f"DEBUG: [actualizar_ot_en_excel] ID actual en celda ({found_row},{id_col_idx}): '{current_id_in_excel_cell}' (Tipo: {type(current_id_in_excel_cell)})")
        
        # Ensure the ID is written back as an integer
        try:
            id_to_write = int(str(current_id_in_excel_cell).strip()) if current_id_in_excel_cell is not None else int(ot_id)
        except (ValueError, TypeError):
            id_to_write = int(ot_id) # Fallback to the ID passed to the function
            print(f"Advertencia: [actualizar_ot_en_excel] No se pudo convertir el ID actual de la celda a entero. Usando ot_id: {id_to_write}")

        ws.cell(row=found_row, column=id_col_idx, value=id_to_write)
        print(f"DEBUG: [actualizar_ot_en_excel] ID {id_to_write} escrito explícitamente en celda ({found_row},{id_col_idx}).")


        for col_name, col_idx_map in EXCEL_COLUMN_MAP.items(): 
            actual_col_idx = header_to_col_idx.get(col_name) 
            
            # Skip ID column as we handled it explicitly above
            if col_name == 'ID':
                continue

            if actual_col_idx is not None and actual_col_idx != -1: 
                if col_name in updated_data: 
                    value_to_write = updated_data.get(col_name)
                    if col_name in ['Visto Bueno', 'Dato Empresa', 'Factura', 'ID Empresa', 'Solicitud O/D', 'Solicitud Anexo']:
                        value_to_write = "SI" if value_to_write == 'SI' or value_to_write == True else "NO"
                    elif value_to_write is None:
                        value_to_write = ''
                    ws.cell(row=found_row, column=actual_col_idx, value=value_to_write)
                elif updated_data.get(col_name) is None: 
                    if col_name not in ['Visto Bueno', 'Dato Empresa', 'Factura', 'ID Empresa', 'Solicitud O/D', 'Solicitud Anexo']:
                        ws.cell(row=found_row, column=actual_col_idx, value='')
            else:
                print(f"Advertencia: Columna '{col_name}' no encontrada en el Excel para actualizar.")

        _recreate_excel_table(ws)

        wb.save(ARCHIVO_LIBRO_RADIO_EXCEL)
        print(f"DEBUG: [actualizar_ot_en_excel] OT {ot_id} actualizada exitosamente en: {ARCHIVO_LIBRO_RADIO_EXCEL}")
        return True
    except Exception as e:
        print(f"ERROR: [actualizar_ot_en_excel] Fallo al actualizar OT en Excel: {e}")
        return False

def obtener_ot_por_id(ot_id):
    """
    Busca una OT específica por su ID en el archivo Excel.
    Devuelve un diccionario con los datos de la OT o None si no se encuentra.
    """
    print(f"DEBUG: [obtener_ot_por_id] Buscando OT con ID: {ot_id}")
    all_ots = leer_libro_radio_excel() 
    for ot in all_ots:
        if ot.get('ID') is not None and int(ot.get('ID')) == int(ot_id):
            print(f"DEBUG: [obtener_ot_por_id] OT {ot_id} encontrada.")
            return ot
    print(f"DEBUG: [obtener_ot_por_id] OT {ot_id} no encontrada.")
    return None

def enviar_archivo_para_descarga(file_path):
    """Envía un archivo para descarga."""
    return send_file(file_path, as_attachment=True)


def _ensure_face_dirs():
    Path(FACE_DATA_DIR).mkdir(parents=True, exist_ok=True)


def _write_face_txt_copy(wb):
    lines = []
    sheets_config = [
        (FACE_REGISTRY_SHEET, FACE_HEADERS),
        (FACE_LOG_SHEET, FACE_LOG_HEADERS),
    ]

    for sheet_name, headers in sheets_config:
        if sheet_name not in wb.sheetnames:
            continue

        ws = wb[sheet_name]
        lines.append(f"[{sheet_name}]")
        lines.append("\t".join(headers))

        for row in ws.iter_rows(min_row=2, max_col=len(headers), values_only=True):
            if not row or all(value is None or str(value).strip() == "" for value in row):
                continue
            lines.append("\t".join("" if value is None else str(value) for value in row))

        lines.append("")

    with open(FACE_TXT_FILE, 'w', encoding='utf-8') as txt_file:
        txt_file.write("\n".join(lines).strip() + "\n")


def _save_face_excel(wb):
    try:
        wb.save(FACE_EXCEL_FILE)
        _write_face_txt_copy(wb)
    except PermissionError as e:
        raise PermissionError(
            f"No se pudo guardar '{FACE_EXCEL_FILE}'. Cierra el archivo Excel si está abierto y vuelve a intentar."
        ) from e


def _ensure_face_excel_sheet(sheet_name, headers, table_name):
    from openpyxl import Workbook

    changed = False

    if os.path.exists(FACE_EXCEL_FILE):
        wb = load_workbook(FACE_EXCEL_FILE)
    else:
        wb = Workbook()
        changed = True

    if sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
    else:
        if len(wb.sheetnames) == 1 and wb.active.max_row == 1 and wb.active.max_column == 1 and wb.active['A1'].value is None:
            ws = wb.active
            ws.title = sheet_name
            changed = True
        else:
            ws = wb.create_sheet(title=sheet_name)
            changed = True

    for index, header in enumerate(headers, start=1):
        current = ws.cell(row=1, column=index).value
        if current != header:
            ws.cell(row=1, column=index, value=header)
            changed = True

    if table_name not in ws.tables:
        last_col = get_column_letter(len(headers))
        last_row = max(1, ws.max_row)
        table = Table(displayName=table_name, ref=f"A1:{last_col}{last_row}")
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium9",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )
        ws.add_table(table)
        changed = True

    if changed:
        _save_face_excel(wb)


def _refresh_simple_table(ws, headers, table_name):
    if table_name in ws.tables:
        del ws.tables[table_name]

    last_col = get_column_letter(len(headers))
    last_row = max(1, ws.max_row)
    table = Table(displayName=table_name, ref=f"A1:{last_col}{last_row}")
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium9",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    ws.add_table(table)


def decode_base64_image(data_url):
    if not data_url or ',' not in data_url:
        return None
    try:
        _, encoded = data_url.split(',', 1)
        binary_data = base64.b64decode(encoded)
        nparr = np.frombuffer(binary_data, np.uint8)
        return cv2.imdecode(nparr, cv2.IMREAD_COLOR)
    except Exception:
        return None


def _extract_face(frame):
    if frame is None or frame.size == 0:
        return None

    gray = cv2.cvtColor(frame, cv2.COLOR_BGR2GRAY)
    face_classifier = cv2.CascadeClassifier(CASCADE_PATH)
    if face_classifier.empty():
        return None

    faces = face_classifier.detectMultiScale(gray, scaleFactor=1.2, minNeighbors=5, minSize=(70, 70))
    if len(faces) == 0:
        return None

    x, y, w, h = max(faces, key=lambda f: f[2] * f[3])
    face = gray[y:y + h, x:x + w]
    return cv2.resize(face, (200, 200))


def _safe_name(name):
    return "".join(ch for ch in name.strip() if ch.isalnum() or ch in ('_', '-', ' ')).strip()


def register_face_from_frame(person_name, frame):
    _ensure_face_dirs()
    _ensure_face_excel_sheet(FACE_REGISTRY_SHEET, FACE_HEADERS, FACE_REGISTRY_TABLE)

    clean_name = _safe_name(person_name)
    if not clean_name:
        return False, "Nombre inválido."

    face = _extract_face(frame)
    if face is None:
        return False, "No se detectó un rostro claro en la imagen."

    filename = f"{clean_name}_{uuid.uuid4().hex[:8]}.png"
    relative_path = os.path.join("faces", filename)
    absolute_path = os.path.join(DATA_DIR, relative_path)
    cv2.imwrite(absolute_path, face)

    wb = load_workbook(FACE_EXCEL_FILE)
    ws = wb[FACE_REGISTRY_SHEET]
    next_id = ws.max_row
    now = datetime.now()
    ws.append([
        next_id,
        clean_name,
        relative_path,
        now.strftime("%Y-%m-%d %H:%M:%S"),
    ])
    _refresh_simple_table(ws, FACE_HEADERS, FACE_REGISTRY_TABLE)
    _save_face_excel(wb)
    return True, f"Rostro de {clean_name} registrado correctamente."


def _load_registered_faces():
    _ensure_face_excel_sheet(FACE_REGISTRY_SHEET, FACE_HEADERS, FACE_REGISTRY_TABLE)
    wb = load_workbook(FACE_EXCEL_FILE)
    ws = wb[FACE_REGISTRY_SHEET]
    entries = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or not row[1] or not row[2]:
            continue
        name = str(row[1]).strip()
        rel_path = str(row[2]).strip()
        full_path = os.path.join(DATA_DIR, rel_path)
        if not os.path.exists(full_path):
            continue
        img = cv2.imread(full_path, cv2.IMREAD_GRAYSCALE)
        if img is None:
            continue
        img = cv2.resize(img, (200, 200))
        entries.append((name, img))
    return entries


def _best_face_match(face_gray, candidates):
    best_name = None
    best_score = -1.0
    for name, candidate_img in candidates:
        score = cv2.compareHist(
            cv2.calcHist([face_gray], [0], None, [256], [0, 256]),
            cv2.calcHist([candidate_img], [0], None, [256], [0, 256]),
            cv2.HISTCMP_CORREL,
        )
        if score > best_score:
            best_score = score
            best_name = name
    return best_name, float(best_score)


def _next_access_type_for_person(person_name):
    _ensure_face_excel_sheet(FACE_LOG_SHEET, FACE_LOG_HEADERS, FACE_LOG_TABLE)
    wb = load_workbook(FACE_EXCEL_FILE)
    ws = wb[FACE_LOG_SHEET]
    last_type = None
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row and row[1] == person_name:
            last_type = row[2]
    return "SALIDA" if last_type == "INGRESO" else "INGRESO"


def register_access_event(person_name, event_type):
    _ensure_face_excel_sheet(FACE_LOG_SHEET, FACE_LOG_HEADERS, FACE_LOG_TABLE)
    wb = load_workbook(FACE_EXCEL_FILE)
    ws = wb[FACE_LOG_SHEET]
    next_id = ws.max_row
    now = datetime.now()
    ws.append([
        next_id,
        person_name,
        event_type,
        now.strftime("%Y-%m-%d"),
        now.strftime("%H:%M:%S"),
        now.strftime("%Y-%m-%d %H:%M:%S"),
    ])
    _refresh_simple_table(ws, FACE_LOG_HEADERS, FACE_LOG_TABLE)
    _save_face_excel(wb)
    return {
        "nombre": person_name,
        "tipo": event_type,
        "fecha": now.strftime("%Y-%m-%d"),
        "hora": now.strftime("%H:%M:%S"),
        "timestamp": now.strftime("%Y-%m-%d %H:%M:%S"),
    }


def list_face_access_events(limit=None):
    _ensure_face_excel_sheet(FACE_LOG_SHEET, FACE_LOG_HEADERS, FACE_LOG_TABLE)
    wb = load_workbook(FACE_EXCEL_FILE)
    ws = wb[FACE_LOG_SHEET]

    events = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row:
            continue

        row_id = row[0]
        nombre = row[1]
        tipo = row[2]
        fecha = row[3]
        hora = row[4]
        timestamp = row[5]

        if not nombre or not tipo:
            continue

        events.append({
            "id": row_id,
            "nombre": str(nombre),
            "tipo": str(tipo),
            "fecha": str(fecha) if fecha is not None else "",
            "hora": str(hora) if hora is not None else "",
            "timestamp": str(timestamp) if timestamp is not None else "",
        })

    def _event_sort_key(event):
        raw_ts = event.get("timestamp") or ""
        try:
            return datetime.strptime(raw_ts, "%Y-%m-%d %H:%M:%S")
        except Exception:
            return datetime.min

    events.sort(key=_event_sort_key, reverse=True)

    if isinstance(limit, int) and limit > 0:
        return events[:limit]
    return events


def scan_face_and_register_access(frame):
    face = _extract_face(frame)
    if face is None:
        return False, "No se detectó un rostro en la imagen.", None

    candidates = _load_registered_faces()
    if not candidates:
        return False, "No hay rostros registrados todavía.", None

    matched_name, score = _best_face_match(face, candidates)
    if matched_name is None or score < 0.72:
        return False, "Rostro no reconocido.", None

    next_type = _next_access_type_for_person(matched_name)
    event_data = register_access_event(matched_name, next_type)
    return True, f"{matched_name} registrado: {next_type}", event_data
